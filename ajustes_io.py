# -*- coding: utf-8 -*-
"""
Aplicacao de uma planilha de AJUSTES (patch esparso) sobre a planilha migrada
de uma conta, chaveada por NUMERO DO PROCESSO.

E um fluxo diferente do resto do padronizador. Nos outros, a planilha e a
FONTE e o app decide o que mudar. Aqui chega uma planilha que JA e a lista de
mudancas: 22 colunas do padrao ADVBOX, so com as celulas que precisam mudar
preenchidas. Celula em branco e uma INSTRUCAO ("esse dado migrado ja esta
certo, nao mexa"), nao falta de dado - inverter isso apagaria dado bom em
massa, por isso e a Regra 1 do proprio arquivo.

Regras, tiradas da aba LEIA-ME do arquivo de ajustes:
  1. campo vazio na correcao nao gera alteracao;
  2. datas em DD/MM/AAAA; valor "0" tratado como vazio;
  3. nomes ja vem resolvidos no arquivo (a comparacao por lista/acento/ordem
     foi feita na geracao dele, nao aqui);
  4. variacao de grafia do mesmo cliente nao entra como troca de nome.

A aba "Sem nome na CORRECAO" lista os processos em que a base correta traz
"NAO INFORMADO"/vazio no nome: o nome migrado e de OUTRO cliente, mas nao ha
nome certo pra por no lugar. Sao os MESMOS processos da aba principal (um
subconjunto), repetidos so pra expor essa decisao. Por padrao nao se mexe no
nome - quem decide limpar e o usuario, item a item.
"""
import os
import re

import openpyxl

# As 22 colunas do padrao ADVBOX, na ordem em que aparecem nos dois arquivos.
COLUNAS_ADVBOX = (
    "NOME DO CLIENTE",
    "PARTE CONTRÁRIA",
    "TIPO DE AÇÃO",
    "GRUPO DE AÇÃO",
    "FASE PROCESSUAL",
    "ETAPA",
    "NÚMERO DO PROCESSO",
    "PROCESSO ORIGINÁRIO",
    "TRIBUNAL",
    "VARA",
    "COMARCA",
    "PROTOCOLO",
    "EXPECTATIVA/VALOR DA CAUSA",
    "VALOR HONORÁRIOS",
    "PASTA",
    "DATA CADASTRO",
    "DATA FECHAMENTO",
    "DATA TRANSITO",
    "DATA ARQUIVAMENTO",
    "DATA REQUERIMENTO",
    "RESPONSÁVEL",
    "ANOTAÇÕES GERAIS",
)

COL_PROCESSO = "NÚMERO DO PROCESSO"
COL_NOME = "NOME DO CLIENTE"

ABA_AJUSTES = "Processos"
ABA_SEM_NOME = "Sem nome na CORREÇÃO"

_SO_ALNUM = re.compile(r"[^0-9A-Za-z]")


def normalizar_processo(numero):
    """Chave de casamento entre os dois arquivos.

    Tira pontuacao do CNJ pra o mesmo processo casar mesmo se um arquivo
    trouxer "1234567-89.2021.8.16.0001" e o outro so os digitos.
    """
    if numero is None:
        return ""
    return _SO_ALNUM.sub("", str(numero)).upper()


def _vazio(valor):
    """Regra 2: alem de None/string vazia, o literal "0" conta como vazio.

    Sem isso, um "0" vindo de campo numerico nao preenchido zeraria um valor
    de honorarios ou uma data que estavam corretos na planilha migrada.
    """
    if valor is None:
        return True
    texto = str(valor).strip()
    return texto == "" or texto == "0"


def _texto(valor):
    return "" if valor is None else str(valor).strip()


def _ler_aba(ws):
    """Devolve {chave_processo: {coluna: valor}} so com as celulas preenchidas.

    Ignora colunas fora do padrao de 22 (se o arquivo vier com extras) e
    linhas sem numero de processo - sem a chave nao ha como aplicar nada.
    """
    cabecalho = [c.value for c in ws[1]]
    try:
        idx_proc = cabecalho.index(COL_PROCESSO)
    except ValueError:
        raise ValueError(
            f'A aba "{ws.title}" nao tem a coluna "{COL_PROCESSO}" - '
            "sem ela nao da pra saber a qual processo cada ajuste pertence."
        )

    fora = {}
    for linha in range(2, ws.max_row + 1):
        numero = ws.cell(row=linha, column=idx_proc + 1).value
        chave = normalizar_processo(numero)
        if not chave:
            continue
        campos = {}
        for i, nome_col in enumerate(cabecalho):
            if nome_col not in COLUNAS_ADVBOX or nome_col == COL_PROCESSO:
                continue
            valor = ws.cell(row=linha, column=i + 1).value
            if not _vazio(valor):
                campos[nome_col] = _texto(valor)
        fora[chave] = campos
    return fora


def ler_ajustes(path):
    """Le o arquivo de ajustes.

    Retorna {"processos": {chave: {col: val}}, "sem_nome": {chave: {col: val}},
    "n_processos": int, "n_sem_nome": int}.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if ABA_AJUSTES not in wb.sheetnames:
        raise ValueError(
            f'Nao encontrei a aba "{ABA_AJUSTES}" neste arquivo. '
            f"Abas presentes: {', '.join(wb.sheetnames)}"
        )
    processos = _ler_aba(wb[ABA_AJUSTES])
    sem_nome = _ler_aba(wb[ABA_SEM_NOME]) if ABA_SEM_NOME in wb.sheetnames else {}
    wb.close()
    return {
        "processos": processos,
        "sem_nome": sem_nome,
        "n_processos": len(processos),
        "n_sem_nome": len(sem_nome),
    }


def _achar_aba_migrada(wb):
    """A aba de processos da planilha migrada.

    Procura pelo nome usual primeiro e, se nao achar, pela primeira aba que
    tenha as colunas de nome e de numero do processo - o arquivo real tambem
    carrega uma aba "CORRECAO" (a base enriquecida do Escavador, com 32
    colunas) que NAO pode ser confundida com a de destino.
    """
    for nome in wb.sheetnames:
        if nome.strip().lower() == "processos migrados":
            return wb[nome]
    for nome in wb.sheetnames:
        ws = wb[nome]
        cabecalho = {c.value for c in ws[1]}
        if COL_PROCESSO in cabecalho and COL_NOME in cabecalho:
            return ws
    raise ValueError(
        "Nao encontrei nenhuma aba com as colunas "
        f'"{COL_NOME}" e "{COL_PROCESSO}" nesta planilha.'
    )


def aplicar_ajustes(migrados_path, ajustes, out_path, decisoes_sem_nome=None,
                    decisoes_conflito=None, dry_run=False):
    """Aplica os ajustes sobre a planilha migrada.

    `decisoes_sem_nome`: {chave_processo: "limpar"|"manter"} - so "limpar"
    age; qualquer outra coisa (inclusive ausencia) deixa o nome migrado como
    esta, que e o comportamento conservador.

    `decisoes_conflito`: {chave_processo: numero_da_linha} - pra numero de
    processo que aparece em MAIS DE UMA linha do migrado, diz em qual linha
    aplicar. Sem decisao, o ajuste NAO e aplicado em nenhuma delas (ver
    abaixo).

    Por que chave duplicada nao aplica sozinha: no arquivo real de um escritorio,
    os 20 numeros repetidos sao PROCESSOS DIFERENTES que acabaram com o mesmo
    CNJ (ex: 7654321-98.2017.8.24.0002 esta numa linha como acao de familia
    de um cliente e noutra como violencia domestica de outro - e o ajuste
    correspondente e de um TERCEIRO caso, contrato bancario). Aplicar nas
    duas sobrescreveria dois processos reais sem relacao com o dado de um
    terceiro. Sao poucos casos e so uma pessoa consegue dizer qual linha e a
    certa - por isso ficam de fora e sao reportados.

    Retorna {"n_alteracoes", "alteracoes": [...], "sem_correspondencia": [...],
    "conflitos_de_chave": {chave: [linhas]}, "filename"}.
    """
    decisoes_sem_nome = decisoes_sem_nome or {}
    decisoes_conflito = decisoes_conflito or {}
    patches = ajustes["processos"]

    wb = openpyxl.load_workbook(migrados_path)
    ws = _achar_aba_migrada(wb)
    cabecalho = [c.value for c in ws[1]]
    col_por_nome = {nome: i + 1 for i, nome in enumerate(cabecalho) if nome}
    idx_proc = col_por_nome.get(COL_PROCESSO)
    if not idx_proc:
        raise ValueError(f'A planilha migrada nao tem a coluna "{COL_PROCESSO}".')

    # mapeia chave -> linhas (lista, porque o migrado real repete 20 numeros)
    linhas_por_chave = {}
    for linha in range(2, ws.max_row + 1):
        chave = normalizar_processo(ws.cell(row=linha, column=idx_proc).value)
        if chave:
            linhas_por_chave.setdefault(chave, []).append(linha)

    duplicados = {k: v for k, v in linhas_por_chave.items() if len(v) > 1}

    alteracoes = []
    sem_correspondencia = []
    conflitos = {}

    for chave, campos in patches.items():
        linhas = linhas_por_chave.get(chave)
        if not linhas:
            sem_correspondencia.append(chave)
            continue
        if len(linhas) > 1:
            escolhida = decisoes_conflito.get(chave)
            if escolhida in linhas:
                linhas = [escolhida]
            else:
                # sem decisao humana, nao mexe em nenhuma - ver docstring
                conflitos[chave] = list(linhas)
                continue
        for linha in linhas:
            for nome_col, novo in campos.items():
                col = col_por_nome.get(nome_col)
                if not col:
                    continue  # coluna nao existe no destino - ignora em silencio
                atual = _texto(ws.cell(row=linha, column=col).value)
                if atual == novo:
                    continue  # ja esta certo, nao conta como alteracao
                alteracoes.append({
                    "processo": chave,
                    "linha": linha,
                    "coluna": nome_col,
                    "antes": atual,
                    "depois": novo,
                })
                if not dry_run:
                    ws.cell(row=linha, column=col).value = novo

    # aba "Sem nome": so age quando o usuario pediu pra limpar
    col_nome = col_por_nome.get(COL_NOME)
    for chave in ajustes.get("sem_nome", {}):
        if decisoes_sem_nome.get(chave) != "limpar":
            continue
        linhas = linhas_por_chave.get(chave, [])
        if len(linhas) > 1:
            escolhida = decisoes_conflito.get(chave)
            if escolhida not in linhas:
                conflitos.setdefault(chave, list(linhas))
                continue
            linhas = [escolhida]
        for linha in linhas:
            atual = _texto(ws.cell(row=linha, column=col_nome).value)
            if not atual:
                continue
            alteracoes.append({
                "processo": chave,
                "linha": linha,
                "coluna": COL_NOME,
                "antes": atual,
                "depois": "",
            })
            if not dry_run:
                ws.cell(row=linha, column=col_nome).value = None

    if not dry_run:
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        wb.save(out_path)
    wb.close()

    return {
        "n_alteracoes": len(alteracoes),
        "alteracoes": alteracoes,
        "sem_correspondencia": sem_correspondencia,
        # numeros repetidos no migrado que TEM ajuste e ficaram sem aplicar
        "conflitos_de_chave": conflitos,
        # todos os numeros repetidos, tendo ajuste ou nao (diagnostico)
        "duplicados_no_migrado": duplicados,
        "filename": os.path.basename(out_path),
    }
