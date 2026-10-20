# -*- coding: utf-8 -*-
"""Aplicacao de uma planilha de AJUSTES (patch esparso) sobre a planilha
migrada, chaveada por NUMERO DO PROCESSO.

Formato do arquivo de ajustes (ver aba LEIA-ME do proprio arquivo):
  - 22 colunas do padrao ADVBOX, mesma ordem da planilha migrada;
  - so vem preenchido o que precisa mudar - coluna em branco significa "o
    dado migrado ja esta certo, nao mexer" (Regra 1). E o oposto de um
    arquivo normal: ausencia de valor e uma INSTRUCAO, nao falta de dado;
  - NUMERO DO PROCESSO sempre preenchido, e a chave;
  - valor "0" conta como vazio (Regra 2).
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest

from ajustes_io import (
    COLUNAS_ADVBOX,
    aplicar_ajustes,
    ler_ajustes,
    normalizar_processo,
)


# ---------- chave ----------

def test_normaliza_numero_de_processo():
    """A chave tem que casar mesmo se a formatacao do CNJ variar entre os
    dois arquivos (com/sem pontuacao, espaco sobrando)."""
    assert normalizar_processo("1234567-89.2021.8.16.0001") == "12345678920218160001"
    assert normalizar_processo(" 1234567-89.2021.8.16.0001 ") == "12345678920218160001"
    assert normalizar_processo("12345678920218160001") == "12345678920218160001"
    assert normalizar_processo(None) == ""
    assert normalizar_processo("") == ""


# ---------- fixtures ----------

def _escreve(path, abas):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome)
        for linha in linhas:
            ws.append(linha)
    wb.save(path)


def _migrados(tmp_path, linhas):
    p = tmp_path / "migrados.xlsx"
    _escreve(p, {"processos migrados": [list(COLUNAS_ADVBOX)] + linhas})
    return str(p)


def _ajustes(tmp_path, processos, sem_nome=None):
    p = tmp_path / "ajustes.xlsx"
    abas = {"Processos": [list(COLUNAS_ADVBOX)] + processos}
    if sem_nome is not None:
        abas["Sem nome na CORREÇÃO"] = [list(COLUNAS_ADVBOX)] + sem_nome
    _escreve(p, abas)
    return str(p)


def _linha(**campos):
    """Monta uma linha de 22 colunas preenchendo so o que for passado."""
    idx = {c: i for i, c in enumerate(COLUNAS_ADVBOX)}
    linha = [None] * len(COLUNAS_ADVBOX)
    for k, v in campos.items():
        linha[idx[k]] = v
    return linha


NUM = "1234567-89.2021.8.16.0001"


# ---------- regra 1: campo vazio nao altera ----------

def test_campo_vazio_no_ajuste_nao_apaga_dado_bom(tmp_path):
    mig = _migrados(tmp_path, [
        _linha(**{"NOME DO CLIENTE": "JOAO DA SILVA",
                  "TIPO DE AÇÃO": "COBRANCA",
                  "NÚMERO DO PROCESSO": NUM,
                  "RESPONSÁVEL": "DRA ANA"}),
    ])
    aj = _ajustes(tmp_path, [
        _linha(**{"TIPO DE AÇÃO": "EXECUCAO", "NÚMERO DO PROCESSO": NUM}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out)

    ws = openpyxl.load_workbook(out)["processos migrados"]
    hdr = [c.value for c in ws[1]]
    linha = {hdr[i]: ws.cell(row=2, column=i + 1).value for i in range(len(hdr))}
    assert linha["TIPO DE AÇÃO"] == "EXECUCAO"      # alterado
    assert linha["NOME DO CLIENTE"] == "JOAO DA SILVA"  # intacto
    assert linha["RESPONSÁVEL"] == "DRA ANA"           # intacto
    assert res["n_alteracoes"] == 1


def test_valor_zero_conta_como_vazio(tmp_path):
    """Regra 2: '0' no ajuste nao pode zerar um campo que tem dado bom."""
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "VALOR HONORÁRIOS": "1500,00"}),
    ])
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "VALOR HONORÁRIOS": "0"}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out)
    ws = openpyxl.load_workbook(out)["processos migrados"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("VALOR HONORÁRIOS") + 1
    assert ws.cell(row=2, column=col).value == "1500,00"
    assert res["n_alteracoes"] == 0


def test_nao_conta_alteracao_quando_valor_e_igual(tmp_path):
    """Se o ajuste traz exatamente o que ja esta la, nao e mudanca - senao o
    relatorio infla e o usuario perde a confianca no numero."""
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "COBRANCA"}),
    ])
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "COBRANCA"}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out)
    assert res["n_alteracoes"] == 0


# ---------- chave: sem correspondencia e duplicados ----------

def test_ajuste_sem_processo_correspondente_e_reportado(tmp_path):
    mig = _migrados(tmp_path, [_linha(**{"NÚMERO DO PROCESSO": NUM})])
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": "9999999-99.2099.8.99.9999",
                  "TIPO DE AÇÃO": "X"}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out)
    assert res["n_alteracoes"] == 0
    assert len(res["sem_correspondencia"]) == 1


def test_numero_duplicado_no_migrado_nao_aplica_e_vira_conflito(tmp_path):
    """No arquivo real de um escritorio, os 20 numeros repetidos sao PROCESSOS
    DIFERENTES com o mesmo CNJ (uma linha e acao de familia de um cliente,
    outra e violencia domestica de outro) - e o ajuste correspondente e de um
    TERCEIRO caso. Aplicar nas duas sobrescreveria dois processos reais sem
    relacao. Fica de fora e vai pra decisao humana."""
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "A"}),
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "B"}),
    ])
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "NOVO"}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out)
    ws = openpyxl.load_workbook(out)["processos migrados"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("TIPO DE AÇÃO") + 1
    assert ws.cell(row=2, column=col).value == "A"   # intacto
    assert ws.cell(row=3, column=col).value == "B"   # intacto
    assert res["n_alteracoes"] == 0
    assert res["conflitos_de_chave"][normalizar_processo(NUM)] == [2, 3]


def test_conflito_aplica_na_linha_que_o_usuario_escolheu(tmp_path):
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "A"}),
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "B"}),
    ])
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "NOVO"}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out,
                          decisoes_conflito={normalizar_processo(NUM): 3})
    ws = openpyxl.load_workbook(out)["processos migrados"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("TIPO DE AÇÃO") + 1
    assert ws.cell(row=2, column=col).value == "A"      # nao escolhida, intacta
    assert ws.cell(row=3, column=col).value == "NOVO"   # escolhida
    assert res["n_alteracoes"] == 1
    assert res["conflitos_de_chave"] == {}


# ---------- aba "Sem nome na CORRECAO" ----------

def test_sem_nome_nao_mexe_por_padrao(tmp_path):
    """Por padrao o nome migrado fica como esta - e o comportamento
    conservador da propria aba "Processos", que deixa o nome em branco."""
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "NOME DO CLIENTE": "CLIENTE ERRADO"}),
    ])
    aj = _ajustes(
        tmp_path,
        [_linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "X"})],
        sem_nome=[_linha(**{"NÚMERO DO PROCESSO": NUM, "NOME DO CLIENTE": "NÃO INFORMADO"})],
    )
    out = str(tmp_path / "out.xlsx")
    dados = ler_ajustes(aj)
    assert len(dados["sem_nome"]) == 1
    res = aplicar_ajustes(mig, dados, out)
    ws = openpyxl.load_workbook(out)["processos migrados"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("NOME DO CLIENTE") + 1
    assert ws.cell(row=2, column=col).value == "CLIENTE ERRADO"


def test_sem_nome_limpa_quando_usuario_decide(tmp_path):
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "NOME DO CLIENTE": "CLIENTE ERRADO"}),
    ])
    aj = _ajustes(
        tmp_path,
        [_linha(**{"NÚMERO DO PROCESSO": NUM})],
        sem_nome=[_linha(**{"NÚMERO DO PROCESSO": NUM, "NOME DO CLIENTE": "NÃO INFORMADO"})],
    )
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(
        mig, ler_ajustes(aj), out,
        decisoes_sem_nome={normalizar_processo(NUM): "limpar"},
    )
    ws = openpyxl.load_workbook(out)["processos migrados"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("NOME DO CLIENTE") + 1
    assert ws.cell(row=2, column=col).value in (None, "")
    assert res["n_alteracoes"] == 1


# ---------- leitura ----------

def test_ler_ajustes_ignora_colunas_vazias(tmp_path):
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "X"}),
    ])
    dados = ler_ajustes(aj)
    chave = normalizar_processo(NUM)
    assert set(dados["processos"][chave]) == {"TIPO DE AÇÃO"}


def test_ler_ajustes_sem_aba_sem_nome(tmp_path):
    aj = _ajustes(tmp_path, [_linha(**{"NÚMERO DO PROCESSO": NUM})])
    dados = ler_ajustes(aj)
    assert dados["sem_nome"] == {}


def test_dry_run_nao_escreve_arquivo(tmp_path):
    mig = _migrados(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "A"}),
    ])
    aj = _ajustes(tmp_path, [
        _linha(**{"NÚMERO DO PROCESSO": NUM, "TIPO DE AÇÃO": "B"}),
    ])
    out = str(tmp_path / "out.xlsx")
    res = aplicar_ajustes(mig, ler_ajustes(aj), out, dry_run=True)
    assert res["n_alteracoes"] == 1
    assert res["alteracoes"][0]["antes"] == "A"
    assert res["alteracoes"][0]["depois"] == "B"
    assert not os.path.exists(out)
