# -*- coding: utf-8 -*-
"""Fluxo completo (upload -> analyze -> apply) dos homonimos de TIPO DE ACAO /
GRUPO DE ACAO, via Flask test client - garante que a fiacao em app.py (rotas,
rules_store, escrita na planilha) funciona junto, nao so o calculo isolado em
matching.py (ver tests/test_label_clusters.py)."""
import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest

import app as app_module
import rules_store


@pytest.fixture()
def client(tmp_path, monkeypatch):
    # isola a persistencia de decisoes num arquivo temporario, pra um teste
    # nunca gravar (nem ler) o data/decisions.json de verdade do usuario.
    monkeypatch.setattr(rules_store, "STORE_PATH", str(tmp_path / "decisions.json"))
    monkeypatch.setattr(rules_store, "DATA_DIR", str(tmp_path))
    # a memoria vem desligada por padrao (ver rules_store.MEMORIA); este teste
    # e justamente sobre o ida-e-volta da persistencia, entao liga de proposito
    # - e so dentro do arquivo temporario acima.
    monkeypatch.setattr(rules_store, "MEMORIA", True)
    app_module.app.config["TESTING"] = True
    with app_module.app.test_client() as c:
        yield c


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Processos"
    ws.append(["NOME DO CLIENTE", "PARTE CONTRARIA", "TIPO DE ACAO", "GRUPO DE ACAO"])
    linhas = [
        ("FULANO DA SILVA", "BANCO X", "COBRANCA", "CIVEL"),
        ("BELTRANO SOUZA", "BANCO Y", "ACAO DE COBRANCA", "ACAO CIVEL"),
        ("CICRANO LIMA", "BANCO Z", "DIVORCIO", "FAMILIA"),
        ("FULANA COSTA", "BANCO W", "DIVORCIO LITIGIOSO", "FAMILIA"),
        ("BELTRANA ROCHA", "BANCO V", "DIVORCIO CONSENSUAL", "FAMILIA"),
    ]
    for row in linhas:
        ws.append(list(row))
    wb.save(path)


def test_fluxo_completo_tipo_e_grupo_acao(client, tmp_path):
    xlsx_path = tmp_path / "esc.xlsx"
    _make_xlsx(xlsx_path)

    with open(xlsx_path, "rb") as f:
        resp = client.post(
            "/api/upload",
            data={"escavador": (io.BytesIO(f.read()), "esc.xlsx")},
            content_type="multipart/form-data",
        )
    assert resp.status_code == 200, resp.get_data(as_text=True)
    session_id = resp.get_json()["session_id"]

    resp = client.post("/api/analyze", json={"session_id": session_id, "threshold": 88})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    data = resp.get_json()

    tipo_clusters = data["tipo_acao_clusters"]
    grupo_clusters = data["grupo_acao_clusters"]

    # COBRANCA + ACAO DE COBRANCA -> 1 cluster de 2
    cobranca = [c for c in tipo_clusters if len(c["members"]) == 2
                and any("COBRANCA" in v["text"].upper() for v in c["variants"])]
    assert len(cobranca) == 1

    # DIVORCIO/LITIGIOSO/CONSENSUAL: LITIGIOSO e CONSENSUAL sao qualificadores
    # de verdade (tipos opostos de divorcio), nao conectores - com a regra de
    # igualdade de conteudo, nenhum dos tres agrupa com os outros dois.
    divorcio = [c for c in tipo_clusters if any("DIVORCIO" in v["text"].upper() for v in c["variants"])]
    assert divorcio == []

    # CIVEL + ACAO CIVEL em GRUPO DE ACAO ("ACAO" e conector, some do conteudo)
    civel = [c for c in grupo_clusters if any("CIVEL" in v["text"].upper() for v in c["variants"])]
    assert len(civel) == 1

    # aplica: COBRANCA -> "ACAO DE COBRANCA" (grupo inteiro), decisoes com touched=True
    cobranca_cluster = cobranca[0]
    tipo_decisions = [{
        "cluster_id": cobranca_cluster["id"],
        "members": cobranca_cluster["members"],
        "excluded_members": [],
        "action": "merge",
        "canonical": "ACAO DE COBRANCA",
        "touched": True,
    }]
    resp = client.post("/api/apply", json={
        "session_id": session_id,
        "decisions": [],
        "tipo_acao_cluster_decisions": tipo_decisions,
        "persist_rules": True,
        "dry_run": False,
    })
    assert resp.status_code == 200, resp.get_data(as_text=True)
    result = resp.get_json()["result"]
    assert result["escavador"]["n_changes"] >= 1

    # confere na planilha de saida: a linha 2 (COBRANCA) virou ACAO DE COBRANCA
    out_path = app_module.SESSIONS[session_id]["esc_path"].replace("uploads", "outputs")
    # caminho exato: pega do proprio resultado
    out_filename = result["escavador"]["filename"]
    session_dir = os.path.dirname(os.path.dirname(app_module.SESSIONS[session_id]["esc_path"]))
    out_full = os.path.join(session_dir, "outputs", out_filename)
    wb2 = openpyxl.load_workbook(out_full)
    ws2 = wb2["Processos"]
    tipo_col = [c.value for c in ws2[1]].index("TIPO DE ACAO") + 1
    valores = [ws2.cell(row=r, column=tipo_col).value for r in range(2, ws2.max_row + 1)]
    assert "ACAO DE COBRANCA" in valores
    assert "COBRANCA" not in valores  # a variante curta foi substituida

    # a decisao "touched" tem que ter sido persistida no namespace certo
    bl, cf = rules_store.load_label_cluster_sets("tipo_acao")
    assert len(cf) >= 1
    bl_g, cf_g = rules_store.load_label_cluster_sets("grupo_acao")
    assert cf_g == {}  # nao vazou pro outro campo
