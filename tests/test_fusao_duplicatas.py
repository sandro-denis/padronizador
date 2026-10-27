# -*- coding: utf-8 -*-
"""Resolver duplicata nao pode apagar dado em silencio.

Antes, marcar uma linha duplicada como "remover" simplesmente apagava a linha.
Se o telefone, o RG ou o endereco so existiam ali - que e justamente o motivo
comum de alguem ter sido cadastrado duas vezes - esse dado ia junto, sem
aparecer em lugar nenhum.

Agora, campo a campo:
  - vazio na linha que fica -> preenche (mesma regra do ADVBOX: nunca
    sobrescreve dado que ja existia);
  - igual -> nao faz nada;
  - DIFERENTE -> nao sobrescreve e nao descarta: vira texto em ANOTACOES
    GERAIS, dizendo de qual coluna veio.

O terceiro caso e de proposito o unico que exige leitura humana: dois valores
diferentes pro mesmo campo da mesma pessoa e uma divergencia que so quem
conhece o cliente resolve.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest

import excel_io

QUEBRA = chr(10)


def _planilha(path, com_anotacoes=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CLIENTES"
    cabecalho = ["NOME", "CPF", "CELULAR", "RG"]
    if com_anotacoes:
        cabecalho.append("ANOTAÇÕES GERAIS")
    ws.append(cabecalho)
    fica = ["FULANA DE TAL", "111.111.111-11", "(41) 1111-1111", ""]
    sai = ["FULANA DE TAL", "111.111.111-11", "(41) 9999-9999", "12.345.678-9"]
    if com_anotacoes:
        fica.append("cliente antiga")
        sai.append("prefere ser chamada de Fulaninha")
    ws.append(fica)
    ws.append(sai)
    wb.save(path)
    return path


def _aplicar(tmp_path, com_anotacoes=True, merge_rows=None):
    entrada = _planilha(str(tmp_path / "cli.xlsx"), com_anotacoes)
    dados = excel_io.read_clientes_rows(entrada)
    saida = str(tmp_path / "saida.xlsx")
    excel_io.apply_clientes(
        entrada, dados, {}, saida,
        rows_to_delete={3},
        merge_rows={3: 2} if merge_rows is None else merge_rows,
    )
    wb = openpyxl.load_workbook(saida)
    return wb["CLIENTES"]


def _coluna(ws, titulo):
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(row=1, column=c).value or "").strip().upper() == titulo:
            return c
    return None


def _anotacao(ws):
    col = _coluna(ws, "ANOTAÇÕES GERAIS")
    assert col, "a coluna de anotacoes devia existir"
    return str(ws.cell(row=2, column=col).value or "")


# ---------- campo vazio: preenche ----------

def test_campo_que_so_existia_na_duplicata_e_aproveitado(tmp_path):
    ws = _aplicar(tmp_path)
    assert ws.cell(row=2, column=_coluna(ws, "RG")).value == "12.345.678-9", \
        "o RG so existia na linha removida e tinha que ter vindo pra linha que fica"


def test_nao_sobra_a_linha_removida(tmp_path):
    """O que NAO pode mudar: a duplicata continua saindo da planilha."""
    ws = _aplicar(tmp_path)
    assert ws.max_row == 2, "devia ter sobrado so o cabecalho e uma linha"


# ---------- campo igual: silencio ----------

def test_campo_identico_nao_vira_anotacao(tmp_path):
    """Anotar o que e igual so faria barulho na ficha."""
    assert "CPF" not in _anotacao(tmp_path and _aplicar(tmp_path))


# ---------- campo diferente: anota, nao sobrescreve ----------

def test_valor_divergente_nao_sobrescreve_o_que_ja_estava(tmp_path):
    ws = _aplicar(tmp_path)
    assert ws.cell(row=2, column=_coluna(ws, "CELULAR")).value == "(41) 1111-1111", \
        "o celular da linha que fica nao podia ser trocado pelo da duplicata"


def test_valor_divergente_vai_pra_anotacao_com_o_nome_da_coluna(tmp_path):
    texto = _anotacao(tmp_path and _aplicar(tmp_path))
    assert "CELULAR: (41) 9999-9999" in texto, texto
    assert "linha 3" in texto, "a anotacao tem que dizer de onde veio"


def test_anotacao_preserva_a_que_ja_existia(tmp_path):
    texto = _anotacao(tmp_path and _aplicar(tmp_path))
    assert texto.startswith("cliente antiga"), texto
    assert QUEBRA in texto


def test_anotacao_da_duplicata_tambem_e_preservada(tmp_path):
    texto = _anotacao(tmp_path and _aplicar(tmp_path))
    assert "prefere ser chamada de Fulaninha" in texto


# ---------- planilha sem a coluna ----------

def test_cria_a_coluna_quando_a_planilha_nao_tem(tmp_path):
    """Sem coluna nao ha onde anotar - e nao anotar seria repetir o proprio
    bug que isso conserta. ANOTACOES GERAIS e coluna do padrao ADVBOX."""
    ws = _aplicar(tmp_path, com_anotacoes=False)
    col = _coluna(ws, "ANOTAÇÕES GERAIS")
    assert col, "a coluna tinha que ter sido criada"
    assert "CELULAR: (41) 9999-9999" in str(ws.cell(row=2, column=col).value)


# ---------- sem fusao, nada muda ----------

def test_sem_merge_rows_a_planilha_nao_ganha_anotacao(tmp_path):
    """O que NAO pode mudar: quem nao resolve duplicata nenhuma nao pode ver
    a ficha dos clientes mudar."""
    ws = _aplicar(tmp_path, merge_rows={})
    col = _coluna(ws, "ANOTAÇÕES GERAIS")
    assert str(ws.cell(row=2, column=col).value or "") == "cliente antiga"


# ---------- a fiacao inteira: tela -> app.py -> planilha ----------

def test_fluxo_completo_pela_api(tmp_path, monkeypatch):
    """O `grupo` e o que diz PARA QUAL linha os dados vao. Ele nasce na tela
    (ClientesDuplicatasSection) e so serve se atravessar o app.py inteiro -
    por isso este teste passa pela rota, nao so pelo excel_io."""
    import io as _io
    import app as app_module

    monkeypatch.setattr(app_module, "SESSIONS_DIR", str(tmp_path / "sessions"))
    monkeypatch.setattr(app_module, "SESSIONS", {})
    os.makedirs(str(tmp_path / "sessions"), exist_ok=True)
    app_module.app.config["TESTING"] = True

    entrada = _planilha(str(tmp_path / "cli.xlsx"))
    with open(entrada, "rb") as f:
        dados = f.read()

    with app_module.app.test_client() as c:
        r = c.post("/api/upload",
                   data={"clientes": (_io.BytesIO(dados), "cli.xlsx")},
                   content_type="multipart/form-data")
        assert r.status_code == 200, r.get_data(as_text=True)
        sid = r.get_json()["session_id"]

        r = c.post("/api/apply", json={
            "session_id": sid,
            "dry_run": False,
            "persist_rules": False,
            "clientes_duplicate_decisions": [
                {"row": 2, "action": "manter", "text": "FULANA DE TAL", "grupo": "exact-0"},
                {"row": 3, "action": "remover", "text": "FULANA DE TAL", "grupo": "exact-0"},
            ],
        })
        assert r.status_code == 200, r.get_data(as_text=True)
        nome = r.get_json()["result"]["clientes"]["filename"]

    gerado = os.path.join(app_module.SESSIONS_DIR, sid, "outputs", nome)
    ws = openpyxl.load_workbook(gerado)["CLIENTES"]
    assert ws.max_row == 2, "a duplicata tinha que ter saido"
    assert ws.cell(row=2, column=_coluna(ws, "RG")).value == "12.345.678-9"
    assert "CELULAR: (41) 9999-9999" in _anotacao(ws)
